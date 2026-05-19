#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
save_extract.py — v3.2 (cpg-data-extraction 스킬 저장 스크립트)

1편 논문 추출 결과(JSON)를 받아 단독 엑셀로 저장.
  입력: 99_Scratch/_tmp_<번호>_<study_id>.json
  출력: 90_Output/extracts/AF_extract_<번호>_<study_id>.xlsx

사용법:
    python save_extract.py <input_json_path>

v3.2 변경:
  - 템플릿 sample_v2.4.xlsx → sample_v2.6.xlsx (48열)
  - BASIC_HEADERS 47열 → 48열: ..., rob_other_coi, analysis_set, notes
  - JSON 기본정보에 'analysis_set' 키 (값: ITT / PP / NR) 권장

Exit codes:
    0  성공
    1  JSON 읽기/파싱 오류
    2  검증 실패 (필수 필드 / 타입 / 구조 / 3시트 번호 일치)
    3  파일 쓰기 오류 (디스크/권한)
    4  템플릿 미존재

설계 근거: 99_Scratch/전략C_구축계획.md (Phase 1 결정 ①~⑩)
"""

import sys
import os
import io
import json
import shutil
import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# 버전 태깅 모듈 (같은 scripts/ 폴더)
try:
    from version_info import set_version as _stamp_version
    _STAMP_AVAILABLE = True
except Exception:
    _STAMP_AVAILABLE = False

# ============================================================
# 상수 (Phase 1 결정 반영)
# ============================================================

SCRIPT_DIR    = Path(__file__).resolve().parent      # scripts/
SKILL_DIR     = SCRIPT_DIR.parent                     # cpg-data-extraction/
TEMPLATE_PATH = SKILL_DIR / 'sample_v2.6.xlsx'        # v3.2: 48열 (analysis_set + notes)
SKILL_VERSION = 'v3.2'

# 서식 상수 (⑤ B+)
FONT_NAME        = 'Arial'
FONT_SIZE        = 9
HEADER_FILL_RGB  = '002F4F4F'   # 헤더 (sample과 일치, ARGB prefix '00')
ROB_FILL_RGB     = '00D6EAF8'   # RoB (기본정보 AJ~AT 데이터 행)
EXCLUDE_FILL_RGB = '00FADBD8'   # exclude=Y 행 전체

# 열 인덱스 (1-indexed)
# v3.2: BASIC_HEADERS 기준 AJ(rob_d1_sequence)=36, ..., AT(rob_other_coi)=46, AU(analysis_set)=47, AV(notes)=48
ROB_COL_START = 36  # AJ
ROB_COL_END   = 46  # AT (inclusive)
EXCLUDE_COL   = 10  # J열

# 숫자 서식 '0' 대상
BASIC_NUM_FMT_ZERO   = [1, 4, 18, 19, 21]   # A(번호), D(year), R, S, U
OUTCOME_NUM_FMT_ZERO = [1]                   # A(번호)
HERB_NUM_FMT_ZERO    = [1]                   # A(번호)

# 경고 임계
CELL_MAX_CHARS = 32767   # Excel 한도
FILE_SIZE_MIN  = 3000    # bytes; 3KB 미만 비정상

# ============================================================
# 시트 스키마
# ============================================================

BASIC_HEADERS = [
    '번호','study_id','author','year','title','journal','vol_issue','pages',
    'doi_pmid','exclude','exclude_reason','country','setting','study_design',
    'sample_size','age_E / age_C','sex_E / sex_C','af_type_code','af_type_other',
    'af_type_text','comorbidity_code','comorbidity_text','tcm_pattern',
    'disease_duration','baseline_comparable','comparison_type','intervention_E',
    'intervention_C','intervention_C_method','co_intervention','treatment_period',
    'treatment_sessions','follow_up','funding','outcomes_reported',
    'rob_d1_sequence','rob_d1_concealment','rob_d1_baseline','rob_d2_blinding',
    'rob_d2_analysis','rob_d3_dropouts','rob_d3_reasons','rob_d4_assessor',
    'rob_d5_protocol','rob_d5_outcome','rob_other_coi','analysis_set','notes'
]
assert len(BASIC_HEADERS) == 48, f"기본정보 헤더 48열 필요 (현재 {len(BASIC_HEADERS)})"

# 아웃컴 헤더 (엑셀에는 \n 포함, JSON에서는 \n 없는 단순 키)
OUTCOME_HEADERS_EXCEL = [
    '번호','study_id','outcome_std','outcome_original','data_type','importance',
    'time_point','E_n','E_val1\n(Mean/Event)','E_val2\n(SD/Total)','C_n',
    'C_val1\n(Mean/Event)','C_val2\n(SD/Total)','effect_type','effect_value',
    'ci_lower','ci_upper','p_value','direction','event_direction','notes',
    'val2_type'
]
OUTCOME_KEYS_JSON = [
    '번호','study_id','outcome_std','outcome_original','data_type','importance',
    'time_point','E_n','E_val1','E_val2','C_n',
    'C_val1','C_val2','effect_type','effect_value',
    'ci_lower','ci_upper','p_value','direction','event_direction','notes',
    'val2_type'
]
assert len(OUTCOME_HEADERS_EXCEL) == 22
assert len(OUTCOME_KEYS_JSON) == 22

HERB_HEADERS = [
    '번호','study_id','formula_name_kr','formula_name_cn','formulation',
    'composition','administration','note'
]
assert len(HERB_HEADERS) == 8

BASIC_KEYS_JSON = BASIC_HEADERS  # \n 없음 → 그대로 사용
HERB_KEYS_JSON  = HERB_HEADERS   # \n 없음 → 그대로 사용

# ============================================================
# 유틸
# ============================================================

def stdout_utf8():
    """Windows 콘솔(cp949)에서 한글 출력 대응."""
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

def log(msg):
    try:
        print(msg)
    except UnicodeEncodeError:
        sys.stdout.buffer.write((msg + '\n').encode('utf-8'))
        sys.stdout.buffer.flush()

def find_project_root(start):
    """위로 10단계까지 올라가 90_Output 있는 폴더를 찾음."""
    p = start
    for _ in range(10):
        if (p / '90_Output').exists():
            return p
        if p == p.parent:
            break
        p = p.parent
    raise RuntimeError('프로젝트 루트(90_Output 포함) 미탐지')

def detect_suspicious_chars(text):
    """렌더링 문제 가능성 높은 문자 감지 (PUA, Replacement char)."""
    if not isinstance(text, str):
        return []
    flagged = []
    for ch in text:
        cp = ord(ch)
        if 0xE000 <= cp <= 0xF8FF:
            flagged.append(('PUA', ch))
        elif cp == 0xFFFD:
            flagged.append(('REPLACEMENT', ch))
    return flagged

# ============================================================
# 경고 누적
# ============================================================

class SaveWarnings:
    """경고 사유 4종 누적. 저장은 계속 진행, 로그에 기록 (⑥ A+)."""
    CODES = ('CJK_SUSPICIOUS', 'CELL_TRUNC', 'OVERWRITE', 'FILE_TOO_SMALL')

    def __init__(self):
        self.items = []  # list of (code, detail)

    def add(self, code, detail):
        assert code in self.CODES, f'알 수 없는 경고 코드: {code}'
        self.items.append((code, detail))

    def any(self):
        return bool(self.items)

# ============================================================
# 예외
# ============================================================

class ValidationError(Exception):
    """치명적 검증 실패. 저장 중단."""
    pass

# ============================================================
# JSON 로드
# ============================================================

def load_json(json_path: Path):
    if not json_path.exists():
        log(f"[ERROR] JSON 파일 없음: {json_path}")
        sys.exit(1)
    try:
        with json_path.open('r', encoding='utf-8') as f:
            return json.load(f)
    except json.JSONDecodeError as e:
        log(f"[ERROR] JSON 파싱 실패: {e}")
        sys.exit(1)
    except Exception as e:
        log(f"[ERROR] JSON 읽기 실패: {e}")
        sys.exit(1)

# ============================================================
# 검증 (③B + ④B)
# ============================================================

def _type_name(x):
    return type(x).__name__

def validate(data):
    """
    ③B (중간): 필수 키 + 기본 타입
    ④B (구조): 시트 3개, 행 수(기본정보·한의중재_한약 각 1행), 3시트 번호·study_id 일치
    실패 시 ValidationError.
    성공 시 (번호, study_id) 반환.
    """
    errs = []
    required_sheets = ['기본정보', '아웃컴', '한의중재_한약']
    for s in required_sheets:
        if s not in data:
            errs.append(f"필수 시트 키 누락: {s!r}")
    if errs:
        raise ValidationError('\n'.join(errs))

    basic   = data['기본정보']
    outcome = data['아웃컴']
    herb    = data['한의중재_한약']

    if not isinstance(basic, list):
        errs.append(f"기본정보는 list여야 함 (현재: {_type_name(basic)})")
    if not isinstance(outcome, list):
        errs.append(f"아웃컴은 list여야 함 (현재: {_type_name(outcome)})")
    if not isinstance(herb, list):
        errs.append(f"한의중재_한약은 list여야 함 (현재: {_type_name(herb)})")
    if errs:
        raise ValidationError('\n'.join(errs))

    # 행 수 (④B)
    if len(basic) != 1:
        errs.append(f"기본정보 행 수는 정확히 1이어야 함 (현재: {len(basic)})")
    if len(herb) != 1:
        errs.append(f"한의중재_한약 행 수는 정확히 1이어야 함 (현재: {len(herb)})")
    if errs:
        raise ValidationError('\n'.join(errs))

    b0 = basic[0]
    if not isinstance(b0, dict):
        raise ValidationError(f"기본정보[0]은 dict여야 함 (현재: {_type_name(b0)})")

    # 필수 키 + 타입 (③B)
    for key, typ, typname in [
        ('번호', int, 'int'),
        ('study_id', str, 'str'),
        ('author', str, 'str'),
        ('year', int, 'int'),
    ]:
        if key not in b0:
            errs.append(f"기본정보 필수 키 누락: {key!r}")
        elif not isinstance(b0[key], typ):
            errs.append(f"기본정보 {key!r}는 {typname} 필요 (현재: {_type_name(b0[key])}={b0[key]!r})")

    h0 = herb[0]
    if not isinstance(h0, dict):
        errs.append(f"한의중재_한약[0]은 dict여야 함 (현재: {_type_name(h0)})")
    else:
        for key, typ, typname in [('번호', int, 'int'), ('study_id', str, 'str')]:
            if key not in h0:
                errs.append(f"한의중재_한약[0] {key!r} 누락")
            elif not isinstance(h0[key], typ):
                errs.append(f"한의중재_한약[0] {key!r}는 {typname} 필요")

    for i, row in enumerate(outcome):
        if not isinstance(row, dict):
            errs.append(f"아웃컴[{i}]은 dict여야 함")
            continue
        for key, typ, typname in [('번호', int, 'int'), ('study_id', str, 'str')]:
            if key not in row:
                errs.append(f"아웃컴[{i}] {key!r} 누락")
            elif not isinstance(row[key], typ):
                errs.append(f"아웃컴[{i}] {key!r}는 {typname} 필요")

    if errs:
        raise ValidationError('\n'.join(errs))

    # 3시트 번호/study_id 일치 (④B)
    ref_num = b0['번호']
    ref_sid = b0['study_id']
    for i, row in enumerate(outcome):
        if row.get('번호') != ref_num:
            errs.append(f"시트 간 번호 불일치: 아웃컴[{i}]={row.get('번호')!r} vs 기본정보={ref_num!r}")
        if row.get('study_id') != ref_sid:
            errs.append(f"시트 간 study_id 불일치: 아웃컴[{i}]={row.get('study_id')!r} vs 기본정보={ref_sid!r}")
    if h0.get('번호') != ref_num:
        errs.append(f"시트 간 번호 불일치: 한의중재_한약={h0.get('번호')!r} vs 기본정보={ref_num!r}")
    if h0.get('study_id') != ref_sid:
        errs.append(f"시트 간 study_id 불일치: 한의중재_한약={h0.get('study_id')!r} vs 기본정보={ref_sid!r}")

    if errs:
        raise ValidationError('\n'.join(errs))

    return ref_num, ref_sid

# ============================================================
# 워크북 조작
# ============================================================

def clear_data_rows(ws):
    """헤더 1행만 남기고 삭제."""
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

def apply_cell_format(cell, is_rob=False, is_exclude=False):
    """데이터 셀 공통 서식 (⑤ B+)."""
    cell.font = Font(name=FONT_NAME, size=FONT_SIZE)
    cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    if is_exclude:
        cell.fill = PatternFill(start_color=EXCLUDE_FILL_RGB, end_color=EXCLUDE_FILL_RGB, fill_type='solid')
    elif is_rob:
        cell.fill = PatternFill(start_color=ROB_FILL_RGB, end_color=ROB_FILL_RGB, fill_type='solid')

def check_warnings_for_value(val, sheet_name, col_idx, row_idx, warnings):
    """경고 대상 문자 검사."""
    if not isinstance(val, str):
        return
    col_letter = get_column_letter(col_idx)
    if len(val) > CELL_MAX_CHARS:
        warnings.add('CELL_TRUNC',
                     f'{sheet_name} {col_letter}{row_idx} 길이 {len(val)}자 (Excel 한도 {CELL_MAX_CHARS} 초과)')
    susp = detect_suspicious_chars(val)
    if susp:
        codes = ','.join(set(c for c, _ in susp))
        warnings.add('CJK_SUSPICIOUS',
                     f'{sheet_name} {col_letter}{row_idx} 의심 문자 {len(susp)}건 ({codes})')

def write_basic_sheet(ws, row_data, warnings):
    """기본정보: 1행 쓰기."""
    clear_data_rows(ws)
    r = 2
    is_exclude = (str(row_data.get('exclude', '')).strip().upper() == 'Y')

    for col_idx, key in enumerate(BASIC_KEYS_JSON, start=1):
        val = row_data.get(key, None)
        cell = ws.cell(row=r, column=col_idx, value=val)
        is_rob = (ROB_COL_START <= col_idx <= ROB_COL_END)
        apply_cell_format(cell, is_rob=is_rob, is_exclude=is_exclude)
        if col_idx in BASIC_NUM_FMT_ZERO:
            cell.number_format = '0'
        check_warnings_for_value(val, '기본정보', col_idx, r, warnings)

    ws.freeze_panes = 'A2'   # ⑤ B+ 편의 서식

def write_outcome_sheet(ws, rows, warnings):
    """아웃컴: N행 쓰기."""
    clear_data_rows(ws)
    for i, row_data in enumerate(rows):
        r = 2 + i
        for col_idx, key in enumerate(OUTCOME_KEYS_JSON, start=1):
            val = row_data.get(key, None)
            cell = ws.cell(row=r, column=col_idx, value=val)
            apply_cell_format(cell)
            if col_idx in OUTCOME_NUM_FMT_ZERO:
                cell.number_format = '0'
            check_warnings_for_value(val, '아웃컴', col_idx, r, warnings)
    ws.freeze_panes = 'A2'

def write_herb_sheet(ws, row_data, warnings):
    """한의중재_한약: 1행 쓰기."""
    clear_data_rows(ws)
    r = 2
    for col_idx, key in enumerate(HERB_KEYS_JSON, start=1):
        val = row_data.get(key, None)
        cell = ws.cell(row=r, column=col_idx, value=val)
        apply_cell_format(cell)
        if col_idx in HERB_NUM_FMT_ZERO:
            cell.number_format = '0'
        check_warnings_for_value(val, '한의중재_한약', col_idx, r, warnings)
    ws.freeze_panes = 'A2'

def verify_template_headers(wb):
    """템플릿 헤더가 우리 스키마와 일치하는지 확인 (삭제·변경 감지)."""
    errs = []
    mapping = {
        '기본정보': BASIC_HEADERS,
        '아웃컴': OUTCOME_HEADERS_EXCEL,
        '한의중재_한약': HERB_HEADERS,
    }
    for sn, expected in mapping.items():
        if sn not in wb.sheetnames:
            errs.append(f"템플릿에 시트 없음: {sn}")
            continue
        ws = wb[sn]
        for c, exp_h in enumerate(expected, start=1):
            actual = ws.cell(row=1, column=c).value
            if actual != exp_h:
                errs.append(f"{sn} {get_column_letter(c)}1 헤더 불일치: expected={exp_h!r}, actual={actual!r}")
                break  # 한 시트당 첫 불일치만 보고
    if errs:
        raise ValidationError('템플릿 헤더 검증 실패:\n  ' + '\n  '.join(errs))

# ============================================================
# 경고 로그
# ============================================================

def append_warning_log(log_path: Path, filename: str, warnings: SaveWarnings):
    if not warnings.any():
        return
    log_path.parent.mkdir(parents=True, exist_ok=True)
    ts = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    lines = [f'{ts}  {filename}']
    for code, detail in warnings.items:
        lines.append(f'  [WARN:{code}] {detail}')
    lines.append('')
    try:
        with log_path.open('a', encoding='utf-8') as f:
            f.write('\n'.join(lines))
    except Exception as e:
        log(f'[WARN] 경고 로그 쓰기 실패 (무시): {e}')

# ============================================================
# 메인
# ============================================================

def main():
    stdout_utf8()

    if len(sys.argv) != 2:
        log('사용법: python save_extract.py <input_json_path>')
        sys.exit(1)

    json_path = Path(sys.argv[1]).resolve()
    data = load_json(json_path)

    # 검증
    try:
        num, sid = validate(data)
    except ValidationError as e:
        log('[VALIDATION ERROR]')
        log(str(e))
        sys.exit(2)

    # 경로 설정
    try:
        project_root = find_project_root(SCRIPT_DIR)
    except Exception as e:
        log(f'[ERROR] {e}')
        sys.exit(3)

    output_dir = project_root / '90_Output' / 'extracts'
    output_dir.mkdir(parents=True, exist_ok=True)
    output_filename = f'AF_extract_{num}_{sid}.xlsx'
    output_path = output_dir / output_filename
    warning_log_path = output_dir / '_save_warnings.log'

    warnings = SaveWarnings()

    # 덮어쓰기 경고
    if output_path.exists():
        warnings.add('OVERWRITE', f'{output_filename} 이미 존재 (덮어씀)')

    # 템플릿 존재 확인
    if not TEMPLATE_PATH.exists():
        log(f'[ERROR] 템플릿 미존재: {TEMPLATE_PATH}')
        sys.exit(4)

    # 템플릿 복사
    try:
        shutil.copy2(TEMPLATE_PATH, output_path)
    except Exception as e:
        log(f'[ERROR] 템플릿 복사 실패: {e}')
        sys.exit(3)

    # 워크북 열기
    try:
        wb = openpyxl.load_workbook(output_path)
    except Exception as e:
        log(f'[ERROR] 워크북 열기 실패: {e}')
        try:
            output_path.unlink()
        except Exception:
            pass
        sys.exit(3)

    # 템플릿 헤더 검증
    try:
        verify_template_headers(wb)
    except ValidationError as e:
        log('[VALIDATION ERROR - 템플릿 손상]')
        log(str(e))
        try:
            output_path.unlink()
        except Exception:
            pass
        sys.exit(2)

    # 시트 쓰기
    try:
        write_basic_sheet(wb['기본정보'], data['기본정보'][0], warnings)
        write_outcome_sheet(wb['아웃컴'], data['아웃컴'], warnings)
        write_herb_sheet(wb['한의중재_한약'], data['한의중재_한약'][0], warnings)
    except Exception as e:
        log(f'[ERROR] 시트 쓰기 실패: {e}')
        try:
            output_path.unlink()
        except Exception:
            pass
        sys.exit(3)

    # 버전 태깅 (저장 직전)
    if _STAMP_AVAILABLE:
        try:
            _stamp_version(wb, SKILL_VERSION)
        except Exception as e:
            log(f'  (버전 태깅 실패, 무시): {e}')

    # 저장
    try:
        wb.save(output_path)
    except Exception as e:
        log(f'[ERROR] 저장 실패: {e}')
        sys.exit(3)

    # 사후 파일 크기 확인
    try:
        size = output_path.stat().st_size
        if size < FILE_SIZE_MIN:
            warnings.add('FILE_TOO_SMALL', f'{output_filename} 크기 {size}B (3KB 미만)')
    except Exception:
        size = None

    # 사후 검증 (read_only로 재확인)
    try:
        wb2 = openpyxl.load_workbook(output_path, read_only=True)
        basic_rows = wb2['기본정보'].max_row - 1
        outcome_rows = wb2['아웃컴'].max_row - 1
        herb_rows = wb2['한의중재_한약'].max_row - 1
        wb2.close()
    except Exception:
        basic_rows = outcome_rows = herb_rows = '?'

    # 경고 로그 기록
    append_warning_log(warning_log_path, output_filename, warnings)

    # 성공 보고
    log(f'✓ 저장 완료: {output_filename}')
    log(f'  번호={num}, study_id={sid}')
    log(f'  행 수: 기본정보={basic_rows}, 아웃컴={outcome_rows}, 한의중재_한약={herb_rows}')
    log(f'  파일: {output_path}')
    if size is not None:
        log(f'  크기: {size}B')
    if warnings.any():
        log(f'  경고 {len(warnings.items)}건 → {warning_log_path.name}')
        for code, detail in warnings.items:
            log(f'    [{code}] {detail}')
    else:
        log('  경고 없음')

    # ⑦A 성공 시 임시 JSON 자동 삭제 (99_Scratch 안의 _tmp_ 접두 파일만)
    try:
        if (json_path.name.startswith('_tmp_')
                and '99_Scratch' in str(json_path)):
            json_path.unlink()
            log(f'  임시 JSON 삭제: {json_path.name}')
    except Exception as e:
        log(f'  (임시 JSON 삭제 실패, 무시): {e}')

    sys.exit(0)


if __name__ == '__main__':
    main()
